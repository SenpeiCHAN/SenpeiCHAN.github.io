#!/usr/bin/env python3
import json
import re
import zipfile
from pathlib import Path
from urllib.parse import urlparse
from xml.etree import ElementTree as ET

from docx import Document
from openpyxl import load_workbook


URL_RE = re.compile(r"https?://[^\s<>\"]+")


def clean_url(url: str) -> str:
    return url.strip().rstrip(").,;，。；）】]")


def slug(text: str, fallback: str = "未分类") -> str:
    text = (text or "").strip()
    text = re.sub(r"[\\/:*?\"<>|]+", "_", text)
    text = re.sub(r"\s+", "_", text)
    text = text.strip("._ ")
    return text[:80] or fallback


def classify(text: str, url: str = "") -> str:
    blob = f"{text} {url}".lower()
    rules = [
        ("企业与产品资料", ["宇树", "unitree", "优必选", "ubtech", "傅利叶", "fourier", "figure", "tesla", "boston", "agility", "apptronik", "智元", "agibot", "engineai", "逐际", "limx", "product", "产品"]),
        ("市场与行业报告", ["市场", "行业", "报告", "研报", "forecast", "market", "research", "产业", "白皮书", "赛迪", "亿欧", "艾瑞", "高工", "ggii"]),
        ("论文与技术资料", ["论文", "paper", "arxiv", "ieee", "researchgate", "acm", "doi.org", "技术", "算法", "benchmark"]),
        ("竞品与岗位能力", ["prd", "产品经理", "实习", "岗位", "能力", "job", "intern", "pm", "招聘", "boss", "linkedin", "lagou"]),
        ("政策与标准", ["政策", "标准", "法规", "工信", "国家", "标准化", "gb/t", "iso", "policy", "standard"]),
        ("新闻与媒体", ["新闻", "采访", "发布", "公众号", "36kr", "晚点", "机器之心", "量子位", "media", "news", "article"]),
    ]
    for name, keys in rules:
        if any(k in blob for k in keys):
            return name
    host = urlparse(url).netloc.lower()
    if any(x in host for x in ["youtube", "bilibili", "vimeo"]):
        return "视频与演示"
    return "其他资料"


def extract_xlsx(path: Path):
    wb = load_workbook(path, data_only=False)
    items = []
    for ws in wb.worksheets:
        headers = {}
        for cell in ws[1]:
            if cell.value:
                headers[cell.column] = str(cell.value).strip()
        for row in ws.iter_rows():
            row_values = []
            row_urls = []
            for cell in row:
                value = "" if cell.value is None else str(cell.value)
                row_values.append(value)
                if cell.hyperlink and cell.hyperlink.target:
                    row_urls.append((cell.coordinate, clean_url(cell.hyperlink.target), value))
                for match in URL_RE.findall(value):
                    row_urls.append((cell.coordinate, clean_url(match), value))
            if not row_urls:
                continue
            title_parts = [v for v in row_values if v and not URL_RE.search(v)]
            title = " - ".join(title_parts[:3])
            for coord, url, label in row_urls:
                col_index = ws[coord].column
                col_name = headers.get(col_index, f"列{col_index}")
                context = " | ".join(v for v in row_values if v)
                items.append(
                    {
                        "source_file": str(path),
                        "source_type": "excel",
                        "sheet": ws.title,
                        "location": f"{ws.title}!{coord}",
                        "column": col_name,
                        "title": title or label or urlparse(url).netloc,
                        "context": context,
                        "url": url,
                        "category": classify(context, url),
                    }
                )
    return items


def paragraph_text_with_hyperlinks(docx_path: Path):
    ns = {
        "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
        "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    }
    with zipfile.ZipFile(docx_path) as zf:
        doc_xml = ET.fromstring(zf.read("word/document.xml"))
        rels = {}
        try:
            rel_xml = ET.fromstring(zf.read("word/_rels/document.xml.rels"))
            for rel in rel_xml:
                rid = rel.attrib.get("Id")
                target = rel.attrib.get("Target")
                if rid and target:
                    rels[rid] = target
        except KeyError:
            pass

    paragraphs = []
    for p_idx, p in enumerate(doc_xml.findall(".//w:p", ns), 1):
        parts = []
        urls = []
        for node in p:
            if node.tag.endswith("hyperlink"):
                rid = node.attrib.get(f"{{{ns['r']}}}id")
                if rid in rels:
                    urls.append(rels[rid])
                text = "".join(t.text or "" for t in node.findall(".//w:t", ns))
                if text:
                    parts.append(text)
            else:
                text = "".join(t.text or "" for t in node.findall(".//w:t", ns))
                if text:
                    parts.append(text)
        text = "".join(parts).strip()
        for match in URL_RE.findall(text):
            urls.append(clean_url(match))
        if text or urls:
            paragraphs.append((p_idx, text, [clean_url(u) for u in urls]))
    return paragraphs


def extract_docx(path: Path):
    doc = Document(path)
    items = []
    paragraphs = paragraph_text_with_hyperlinks(path)
    text_by_para = {idx: text for idx, text, _ in paragraphs}
    heading_context = "未命名章节"
    for idx, text, urls in paragraphs:
        paragraph = doc.paragraphs[idx - 1] if idx - 1 < len(doc.paragraphs) else None
        style = paragraph.style.name if paragraph and paragraph.style else ""
        if text and ("Heading" in style or "标题" in style):
            heading_context = text
        context = text or heading_context
        needish = any(k in context for k in ["资料", "链接", "报告", "论文", "案例", "竞品", "岗位", "下载", "补充", "参考"])
        for url in urls:
            items.append(
                {
                    "source_file": str(path),
                    "source_type": "word",
                    "sheet": "",
                    "location": f"段落{idx}",
                    "column": "",
                    "title": heading_context if heading_context != "未命名章节" else (text[:80] or urlparse(url).netloc),
                    "context": context,
                    "url": url,
                    "category": classify(context, url),
                }
            )
        if needish and not urls:
            items.append(
                {
                    "source_file": str(path),
                    "source_type": "word_need",
                    "sheet": "",
                    "location": f"段落{idx}",
                    "column": "",
                    "title": heading_context,
                    "context": context,
                    "url": "",
                    "category": classify(context),
                }
            )
    return items


def main():
    import argparse

    parser = argparse.ArgumentParser()
    parser.add_argument("--xlsx", required=True)
    parser.add_argument("--docx", required=True)
    parser.add_argument("--out", required=True)
    args = parser.parse_args()

    out = Path(args.out)
    out.mkdir(parents=True, exist_ok=True)
    items = extract_xlsx(Path(args.xlsx)) + extract_docx(Path(args.docx))
    seen = set()
    deduped = []
    for item in items:
        key = (item["url"], item["source_type"], item["location"], item["context"])
        if key in seen:
            continue
        seen.add(key)
        item["category_slug"] = slug(item["category"])
        deduped.append(item)
    (out / "sources_manifest.json").write_text(json.dumps(deduped, ensure_ascii=False, indent=2), encoding="utf-8")
    with (out / "sources_manifest.csv").open("w", encoding="utf-8") as f:
        f.write("source_type,location,category,title,url,context\n")
        for item in deduped:
            vals = [item["source_type"], item["location"], item["category"], item["title"], item["url"], item["context"]]
            f.write(",".join('"' + str(v).replace('"', '""') + '"' for v in vals) + "\n")
    print(json.dumps({"items": len(deduped), "with_url": sum(1 for i in deduped if i["url"]), "out": str(out)}, ensure_ascii=False))


if __name__ == "__main__":
    main()
